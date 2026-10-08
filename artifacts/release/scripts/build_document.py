"""Run with Codex bundled document Python, not the training environment."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
doc=Document();section=doc.sections[0]
section.page_height=Cm(29.7);section.page_width=Cm(21)
section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2);section.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=doc.styles[name];s.font.name='Microsoft YaHei';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    s.font.size=Pt(10.5 if name=='Normal' else 22 if name=='Title' else 14)
    s.font.italic=False
    s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.widow_control=True
doc.add_heading('PYNQ Z2 皮肤病灶分割系统',0)
doc.add_paragraph('工程实现与调试说明',style='Subtitle')
doc.add_paragraph('本说明介绍板端分割、HDMI 交互、模型训练、定点接口、硬件构建和调试。以随工程提供的验收状态文件为准，软件参考结果不等于已完成上板测试。')
if (root/'artifacts/evaluation_test.json').exists():
    report=json.loads((root/'artifacts/evaluation_test.json').read_text())
    doc.add_heading('已获得的算法证据',1)
    doc.add_paragraph(f"冻结模型 {report['model_id']}。独立测试集 {report['count']} 张，指标在原图尺寸计算。INT8 来自原生 C++ 核心执行；板上精度一致性仍须用相同模型复核。")
    table=doc.add_table(rows=1,cols=3)
    for cell,text in zip(table.rows[0].cells,['版本','平均 Dice','平均 IoU']):cell.text=text
    for name,label in [('fp32_raw','FP32 原始分割'),('int8_raw','INT8 原始分割'),('int8_post','INT8 加冻结后处理')]:
        cells=table.add_row().cells
        for cell,text in zip(cells,[label,f"{report[name]['dice']:.4f}",f"{report[name]['iou']:.4f}"]):cell.text=text
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');tcPr.append(shade)
            for p in cell.paragraphs:p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    doc.add_paragraph('模型、数据和算法证据均保存在 artifacts 与 record。FPGA 性能、资源、功耗及 HDMI 运行状态须分别查看实际报告，不能从此表推断。')
for line in (root/'docs/项目说明.md').read_text(encoding='utf-8').splitlines():
    if not line or line.startswith('# '):continue
    if line.startswith('## '):
        title=line[3:];doc.add_heading(title,1)
        if title=='HDMI 界面与帧同步' and (root/'artifacts/demo_snapshot.png').exists():
            doc.add_picture(str(root/'artifacts/demo_snapshot.png'),width=Cm(17))
            doc.add_paragraph('图示为开发电脑执行原生 C++ 数值参考后的界面预览，不是板卡实拍。')
    else:doc.add_paragraph(line)
doc.core_properties.title='PYNQ Z2 皮肤病灶分割系统项目说明'
doc.core_properties.subject='架构 定点接口 构建 部署 验证 调试'
doc.core_properties.author='项目小组'
for element in list(doc.styles.element.iter())+list(doc.element.iter()):
    if element.tag==qn('w:pBdr'):
        element.getparent().remove(element)
doc.save(root/'docs/项目说明.docx')
print(root/'docs/项目说明.docx')
